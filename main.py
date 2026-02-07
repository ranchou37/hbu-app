import tkinter as tk
from tkinter import messagebox
from tkinter import ttk
import tkinter.font as tkfont
import os
import re
import chardet
import sys
from docx import Document

# 1. 경로 및 초기 설정
if getattr(sys, 'frozen', False):
    BASE_DIR = os.path.dirname(sys.executable)
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))

bible_directory = os.path.join(BASE_DIR, "성경66권_파이션_자료")
hymn_directory = os.path.join(BASE_DIR, "찬송가-가사TXT")
prayer_file_path = os.path.join(BASE_DIR, "prayer.txt")
kidokmoon_file_path = os.path.join(BASE_DIR, "교독문.docx")
sado_path = os.path.join(BASE_DIR, "사도신경.txt")
ju_path = os.path.join(BASE_DIR, "주기도문.txt")

# 2. 데이터 로딩 및 도우미 함수들
def load_txt_file(path):
    if os.path.exists(path):
        with open(path, 'r', encoding='utf-8') as f: return f.read()
    return ""

def load_docx_paragraphs(path):
    if os.path.exists(path):
        try:
            doc = Document(path)
            return [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        except: return []
    return []

# 성경 데이터 로딩
bible = {}
if os.path.exists(bible_directory):
    for filename in sorted(os.listdir(bible_directory)):
        if filename.endswith(".txt"):
            book_key = filename.replace(".txt", "")
            book_key = re.sub(r"^\d+-\d+", "", book_key)
            book_key = re.sub(r"\d+$", "", book_key).strip()
            file_path = os.path.join(bible_directory, filename)
            with open(file_path, "rb") as f:
                raw_data = f.read()
                encoding = chardet.detect(raw_data)['encoding'] or 'utf-8'
            with open(file_path, "r", encoding=encoding) as file:
                verses = {}
                for line in file.readlines():
                    match = re.match(r"([\uac00-\ud7a3]+)(\d+):(\d+)\s+(.*)", line.strip())
                    if match:
                        chapter, verse, text_line = int(match[2]), int(match[3]), match[4]
                        verses[f"{chapter}:{verse}"] = text_line
                bible[book_key] = verses

dict_book = {
    "창": "창세기", "출": "출애굽기", "레": "레위기", "민": "민수기", "신": "신명기",
    "수": "여호수아", "삿": "사사기", "룻": "룻기", "삼상": "사무엘상", "삼하": "사무엘하",
    "왕상": "열왕기상", "왕하": "열왕기하", "대상": "역대상", "대하": "역대하",
    "스": "에스라", "느": "느헤미야", "에": "에스더", "욥": "욥기", "시": "시편",
    "잠": "잠언", "전": "전도서", "아": "아가", "사": "이사야", "렘": "예레미야",
    "애": "예레미야애가", "겔": "에스겔", "단": "다니엘", "호": "호세아", "욜": "요엘",
    "암": "아모스", "옵": "오바댜", "욘": "요나", "미": "미가", "나": "나훔",
    "합": "하박국", "습": "스바냐", "학": "학개", "슥": "스가랴", "말": "말라기",
    "마": "마태복음", "막": "마가복음", "눅": "누가복음", "요": "요한복음",
    "행": "사도행전", "롬": "로마서", "고전": "고린도전서", "고후": "고린도후서",
    "갈": "갈라디아서", "엡": "에베소서", "빌": "빌립보서", "골": "골로새서",
    "살전": "데살로니가전서", "살후": "데살로니가후서", "딤전": "디모데전서",
    "딤후": "디모데후서", "딛": "디도서", "몬": "빌레몬서", "히": "히브리서",
    "약": "야고보서", "벧전": "베드로전서", "벧후": "베드로후서",
    "요일": "요한일서", "요이": "요한이서", "요삼": "요한삼서", "유": "유다서", "계": "요한계시록"
}
all_books_order = list(dict_book.values())
book_full_to_short = {full: short for short, full in dict_book.items()}

bible_categories = {
    "전체": all_books_order, 
    "구약": all_books_order[:39], 
    "신약": all_books_order[39:],
    "모세오경": ["창세기", "출애굽기", "레위기", "민수기", "신명기"],
    "4복음서": ["마태복음", "마가복음", "누가복음", "요한복음"]
}

# 교독문 로딩
docx_paragraphs = load_docx_paragraphs(kidokmoon_file_path)
kidokmoon_titles, kidokmoon_sections = [], []
current_title, current_content = "", ""
for para in docx_paragraphs:
    if para and para[0].isdigit() and '.' in para:
        if current_title:
            kidokmoon_titles.append(current_title); kidokmoon_sections.append(current_content.strip())
        current_title = para; current_content = ""
    else: current_content += para + "\n"
if current_title:
    kidokmoon_titles.append(current_title); kidokmoon_sections.append(current_content.strip())

# 3. GUI 구성
root = tk.Tk()
root.title("성경 및 찬송가 검색 프로그램")
root.geometry("1250x850")

# 제어 변수
is_worship_mode = False
title_font_size = tk.IntVar(value=70)
font_size = tk.IntVar(value=60)

# --- UI 프레임들 ---
top_frame = tk.Frame(root)
top_frame.pack(pady=5, fill="x")

combined_frame = tk.Frame(root)
combined_frame.pack(pady=5, fill="x")

# 본문 프레임
text_frame = tk.Frame(root)
text_frame.pack(fill="both", expand=True, padx=5, pady=5)

ui_frames = [top_frame, combined_frame] 

# --- 1단 구성 ---
tk.Label(top_frame, text="찬송가:").pack(side="left", padx=2)
hymn_entry = tk.Entry(top_frame, width=8)
hymn_entry.pack(side="left", padx=2)

tk.Label(top_frame, text="성경:").pack(side="left", padx=2)
ref_entry = tk.Entry(top_frame, width=12)
ref_entry.pack(side="left", padx=2)

category_var, book_var = tk.StringVar(value="전체"), tk.StringVar(value="전체")
category_combo = ttk.Combobox(top_frame, textvariable=category_var, state="readonly", width=8, values=list(bible_categories.keys()))
category_combo.pack(side="left", padx=2)
book_combo = ttk.Combobox(top_frame, textvariable=book_var, state="readonly", width=10)
book_combo.pack(side="left", padx=2)

def on_category_selected(event=None):
    books = bible_categories.get(category_var.get(), all_books_order)
    book_combo['values'] = ["전체"] + books
    book_var.set("전체")

def on_book_selected(event=None):
    if book_var.get() != "전체":
        short = book_full_to_short.get(book_var.get(), book_var.get())
        ref_entry.delete(0, tk.END)
        ref_entry.insert(0, f"{short} ")

category_combo.bind("<<ComboboxSelected>>", on_category_selected)
book_combo.bind("<<ComboboxSelected>>", on_book_selected)
on_category_selected()

tk.Button(top_frame, text="사도신경", command=lambda: (text.delete("1.0", tk.END), text.insert(tk.END, load_txt_file(sado_path)))).pack(side="left", padx=2)
tk.Button(top_frame, text="주기도문", command=lambda: (text.delete("1.0", tk.END), text.insert(tk.END, load_txt_file(ju_path)))).pack(side="left", padx=2)

dropdown_var = tk.StringVar(root, value="교독문 선택")
if not kidokmoon_titles: kidokmoon_titles = ["없음"]
dropdown_menu = tk.OptionMenu(top_frame, dropdown_var, *kidokmoon_titles, 
    command=lambda title: (text.delete("1.0", tk.END), text.insert(tk.END, f"{title}\n\n{kidokmoon_sections[kidokmoon_titles.index(title)]}")))
dropdown_menu.pack(side="left", padx=5)

font_ctrl_frame = tk.Frame(top_frame)
font_ctrl_frame.pack(side="left", padx=5)
tk.Scale(font_ctrl_frame, from_=30, to=100, orient="horizontal", variable=title_font_size, width=8, label="제목").pack(side="left")
tk.Scale(font_ctrl_frame, from_=40, to=120, orient="horizontal", variable=font_size, width=8, label="본문").pack(side="left")

# --- 2단 구성 (키워드 + 정보) ---
search_sub = tk.Frame(combined_frame)
search_sub.pack(side="left", padx=10)
tk.Label(search_sub, text="키워드:").pack(side="left")
search_entry = tk.Entry(search_sub, width=15)
search_entry.pack(side="left", padx=2)

def find_text(event=None):
    keyword = search_entry.get().strip()
    if not keyword: return
    text.delete("1.0", tk.END)
    results = []
    target_books = [book_var.get()] if book_var.get() != "전체" else bible_categories.get(category_var.get(), all_books_order)
    for b in target_books:
        if b in bible:
            for vk, vt in bible[b].items():
                if keyword in vt: results.append(f"{b} {vk}  {vt}")
    text.insert(tk.END, "\n".join(results) if results else "결과 없음")
    if results:
        idx = "1.0"
        while True:
            idx = text.search(keyword, idx, nocase=True, stopindex=tk.END)
            if not idx: break
            lastidx = f"{idx}+{len(keyword)}c"
            text.tag_add("highlight", idx, lastidx)
            idx = lastidx
        text.tag_config("highlight", background="yellow", foreground="black")

tk.Button(search_sub, text="키워드검색", command=find_text).pack(side="left", padx=2)
search_entry.bind("<Return>", find_text)

tk.Label(combined_frame, text="|", fg="gray").pack(side="left", padx=5)

info_sub = tk.Frame(combined_frame)
info_sub.pack(side="left", padx=5)
info_entries = []
for label in ["제 몇차", "강사", "말씀", "제목"]:
    tk.Label(info_sub, text=label + ":").pack(side="left", padx=1)
    e = tk.Entry(info_sub, width=10)
    e.pack(side="left", padx=2)
    info_entries.append(e)

def display_info():
    num, spk, vrs, tit = [e.get().strip() or " " for e in info_entries]
    text.delete("1.0", tk.END)
    b_font = tkfont.Font(family="맑은 고딕", size=title_font_size.get(), weight="bold")
    text.tag_configure("bold_title", font=b_font, justify="center")
    text.insert("1.0", f"제 {num} 차\n공주사랑중보기도회\n", "bold_title")
    text.insert(tk.END, f"\n강사: {spk}\n성경: {vrs}\n제목: {tit}", "center")

tk.Button(info_sub, text="정보 표시", command=display_info).pack(side="left", padx=5)

# 기도제목 기능
def save_section(tag):
    content = text.get("1.0", tk.END).strip()
    if not content: return
    old = load_txt_file(prayer_file_path)
    pattern = rf"＃{tag}\n(.*?)\n＃{tag}"
    new_data = f"＃{tag}\n{content}\n＃{tag}"
    if re.search(pattern, old, re.DOTALL):
        new = re.sub(pattern, new_data, old, flags=re.DOTALL)
    else: new = old + "\n" + new_data
    with open(prayer_file_path, "w", encoding="utf-8") as f: f.write(new.strip())
    messagebox.showinfo("저장", f"{tag} 저장됨")

def show_tagged_section(tag):
    raw = load_txt_file(prayer_file_path)
    match = re.search(rf"＃{tag}\n(.*?)\n＃{tag}", raw, re.DOTALL)
    if match:
        text.delete("1.0", tk.END); text.insert(tk.END, match.group(1))
    else: messagebox.showinfo("안내", "내용 없음")

btn_frame = tk.Frame(combined_frame)
btn_frame.pack(side="left", padx=10)
for tag in ["기도제목", "나라", "공주", "회원"]:
    tk.Button(btn_frame, text=tag, command=lambda t=tag: show_tagged_section(t), width=6).pack(side="left", padx=1)
tk.Button(btn_frame, text="저장", command=lambda: save_section("기도제목"), bg="#e1f5fe").pack(side="left", padx=2)

# --- 3단 구성 ---
scrollbar = tk.Scrollbar(text_frame)
scrollbar.pack(side="right", fill="y")
text = tk.Text(text_frame, wrap="word", yscrollcommand=scrollbar.set, font=("맑은 고딕", 60))
text.pack(side="left", fill="both", expand=True)
scrollbar.config(command=text.yview)
text.tag_configure("center", justify="center")

# 성경 검색 로직
def gui_search_reference(event=None):
    ref = ref_entry.get().strip()
    match = re.match(r"([\uac00-\ud7a3]+)\s*(\d+):?(\d+)?(?:-(\d+))?", ref)
    if not match: return
    short, ch, v1, v2 = match.groups()
    book = dict_book.get(short, short)
    if book not in bible: return
    res = []
    ch_int = int(ch)
    if v1 is None:
        v_list = sorted([int(k.split(':')[1]) for k in bible[book].keys() if k.startswith(f"{ch_int}:")])
        for v in v_list: res.append(f"{ch_int}:{v} {bible[book][f'{ch_int}:{v}']}")
    else:
        s = int(v1)
        e = int(v2) if v2 else s
        for v in range(s, e + 1):
            if f"{ch_int}:{v}" in bible[book]: res.append(f"{ch_int}:{v} {bible[book][f'{ch_int}:{v}']}")
    text.insert(tk.END, f"\n\n[{book} {ch_int}장" + (f" {v1}-{v2}절" if v2 else (f" {v1}절" if v1 else "")) + "]\n" + "\n".join(res))
    text.see(tk.END)

def search_hymn(event=None):
    k = hymn_entry.get().strip()
    if not k: return
    results = []
    if os.path.exists(hymn_directory):
        for file in os.listdir(hymn_directory):
            if k in file:
                content = load_txt_file(os.path.join(hymn_directory, file))
                results.append(f"[{file.replace('.txt','')}]\n{content}")
    text.insert(tk.END, "\n\n" + "\n\n".join(results) if results else "\n찬송가 없음")
    text.see(tk.END)

# --- 예배 모드 제어 (수정 완료) ---
def enter_worship_mode(event=None):
    global is_worship_mode
    if is_worship_mode: return
    is_worship_mode = True
    for f in ui_frames: f.pack_forget()
    root.attributes("-fullscreen", True)
    font_size.set(90)

def exit_worship_mode(event=None):
    global is_worship_mode
    if not is_worship_mode: return
    is_worship_mode = False
    root.attributes("-fullscreen", False)
    root.state("zoomed")
    font_size.set(60)
    # UI 다시 배치 (순서 고정)
    text_frame.pack_forget()
    top_frame.pack(pady=5, fill="x")
    combined_frame.pack(pady=5, fill="x")
    text_frame.pack(fill="both", expand=True, padx=5, pady=5)

def update_font_realtime():
    text.config(font=("맑은 고딕", font_size.get()))
    root.after(500, update_font_realtime)

tk.Button(top_frame, text="성경 검색", command=gui_search_reference).pack(side="left", padx=5)
tk.Button(top_frame, text="찬송가 검색", command=search_hymn).pack(side="left", padx=5)
tk.Button(top_frame, text="지우기", command=lambda: text.delete("1.0", tk.END), bg="white").pack(side="left", padx=5)

root.bind("<F1>", enter_worship_mode)
root.bind("<Escape>", exit_worship_mode)
ref_entry.bind("<Return>", gui_search_reference)
hymn_entry.bind("<Return>", search_hymn)

update_font_realtime()
root.mainloop()